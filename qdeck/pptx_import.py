"""Import an existing .pptx (one question per slide) as an editable project skeleton.

Standalone: depends only on python-pptx and Pillow, not on qdeck.deck internals.
See docs/EDITOR_API.md for the project schema this produces.
"""
from __future__ import annotations

import io
import os
import re

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

_RASTER_KEEP = {"png": "png", "jpg": "jpg", "jpeg": "jpg"}
_NOTES_LABEL = re.compile(r"^Q\s*([^|\s]+)")
# '7.' / '7)' / 'Q.43' / 'Q 43.' / 'Q43' / '(12)'
_TEXT_LABEL = re.compile(r"^\s*(Q\s*\.?\s*\d+[.)]?|\(\d+\)|\d+[.)])(?=\s|$)", re.I)
_OPTION = re.compile(r"^\s*(?:\(([a-dA-D1-4])\)|([a-dA-D])\))\s*(.*)$", re.S)
_INLINE_OPT_SPLIT = re.compile(r"\s+(?=\([a-dA-D1-4]\)\s)")


def _iter_shapes(shapes):
    """Yield every leaf shape, descending into groups."""
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(sh.shapes)
        else:
            yield sh


def _picture_image(sh):
    """Return the shape's pptx Image if it is a picture (incl. filled picture placeholder), else None."""
    try:
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE or (getattr(sh, "is_placeholder", False) and hasattr(sh, "image")):
            return sh.image
    except Exception:  # empty picture placeholder, broken rel, ...
        return None
    return None


def _area(sh):
    try:
        return int(sh.width or 0) * int(sh.height or 0)
    except Exception:
        return 0


def _paragraphs(slide):
    """Non-empty text lines from all text frames, in rough reading order (top, then left)."""
    frames = []
    for sh in _iter_shapes(slide.shapes):
        if not getattr(sh, "has_text_frame", False) or not sh.has_text_frame:
            continue
        lines = []
        for p in sh.text_frame.paragraphs:
            for ln in re.split(r"[\n\v\x0b]", p.text or ""):
                ln = ln.strip()
                if ln:
                    lines.append(ln)
        if lines:
            frames.append(((sh.top or 0), (sh.left or 0), lines))
    frames.sort(key=lambda f: (f[0], f[1]))
    return [ln for _, _, lines in frames for ln in lines]


def _notes_first_line(slide):
    try:
        if not slide.has_notes_slide:
            return ""
        tf = slide.notes_slide.notes_text_frame
        text = (tf.text if tf is not None else "") or ""
    except Exception:
        return ""
    for ln in re.split(r"[\r\n\v]", text):
        if ln.strip():
            return ln.strip()
    return ""


def _crop_box(pic):
    def g(name):
        try:
            return float(getattr(pic, name) or 0.0)
        except Exception:
            return 0.0
    x0, y0 = max(0.0, g("crop_left")), max(0.0, g("crop_top"))
    x1, y1 = min(1.0, 1.0 - g("crop_right")), min(1.0, 1.0 - g("crop_bottom"))
    if x1 - x0 <= 0.01 or y1 - y0 <= 0.01:
        return [0.0, 0.0, 1.0, 1.0]
    return [round(x0, 6), round(y0, 6), round(x1, 6), round(y1, 6)]


def _save_image(image, out_dir, seq):
    """Write the picture blob as page_{seq:03d}.{ext}. Returns (path, w, h) or None if unusable."""
    from PIL import Image

    ext = (image.ext or "").lower()
    blob = image.blob
    try:
        with Image.open(io.BytesIO(blob)) as im:
            im.load()
            w, h = im.size
            if ext in _RASTER_KEEP:
                path = os.path.join(out_dir, f"page_{seq:03d}.{_RASTER_KEEP[ext]}")
                with open(path, "wb") as f:
                    f.write(blob)
            elif ext in ("wmf", "emf"):
                return None  # vector metafiles: Pillow cannot rasterise them portably
            else:
                path = os.path.join(out_dir, f"page_{seq:03d}.png")
                conv = im if im.mode in ("RGB", "RGBA", "L", "LA") else im.convert("RGBA")
                conv.save(path, "PNG")
    except Exception:
        return None
    return os.path.abspath(path), w, h


def _number(label):
    m = re.search(r"\d+", label or "")
    return int(m.group()) if m else None


def _split_text(lines):
    """Pull (label, stem, options) out of slide text lines."""
    label = None
    rest = list(lines)
    if rest:
        m = _TEXT_LABEL.match(rest[0])
        if m:
            label = re.sub(r"\s+", "", m.group(1)) if m.group(1).lower().startswith("q") else m.group(1)
            rest[0] = rest[0][m.end():].strip()
    expanded = []
    for ln in rest:
        expanded.extend(s for s in _INLINE_OPT_SPLIT.split(ln) if s.strip())
    stem_lines, options = [], []
    for ln in expanded:
        m = _OPTION.match(ln)
        if m:
            options.append(m.group(3).strip())
        elif options:  # continuation of the previous option
            options[-1] = (options[-1] + " " + ln).strip()
        else:
            stem_lines.append(ln)
    return label, "\n".join(stem_lines).strip(), options


def import_pptx(pptx_path: str, out_dir: str) -> dict:
    """Turn a .pptx into a project dict {pages, items, settings} (see docs/EDITOR_API.md)."""
    os.makedirs(out_dir, exist_ok=True)
    prs = Presentation(pptx_path)
    name = os.path.basename(pptx_path)
    pages, items, used_ids = [], [], set()
    found_blank = False

    for slide_no, slide in enumerate(prs.slides, start=1):
        shapes = list(_iter_shapes(slide.shapes))
        pics = [(sh, img) for sh in shapes for img in [_picture_image(sh)] if img is not None]
        lines = _paragraphs(slide)

        if not pics and not lines:  # blank working slide
            found_blank = True
            if items:
                items[-1]["blank_after"] = True
            continue

        saved = None
        for sh, img in sorted(pics, key=lambda p: _area(p[0]), reverse=True):
            saved = _save_image(img, out_dir, len(pages))
            if saved:
                pic_shape = sh
                break
        if saved is None and not lines:
            continue  # only unusable pictures (wmf/emf/broken): nothing to import

        k = len(items)
        text_label, stem, options = _split_text(lines)
        m = _NOTES_LABEL.match(_notes_first_line(slide))
        label = (m.group(1) if m else None) or text_label or str(k + 1)
        number = _number(label)

        n = number if number is not None and number not in used_ids else None
        if n is None:
            n = k + 1
            while n in used_ids:
                n += 1
        used_ids.add(n)

        item = {"id": f"q{n:04d}", "label": label, "number": number, "lang": "en",
                "render": "snapshot", "stem": "", "options": [], "parts": [], "flags": [],
                "confidence": 1.0, "blank_after": False, "custom_image": False, "rev": 1}
        if saved:
            path, w, h = saved
            seq = len(pages)
            pages.append({"seq": seq, "ref": f"{name}#s{slide_no}", "path": path, "kind": "image", "w": w, "h": h})
            item["parts"] = [{"page": seq, "box": _crop_box(pic_shape)}]
        else:
            item.update(render="text", stem=stem, options=options)
        items.append(item)

    return {"pages": pages, "items": items,
            "settings": {"layout": "fill", "blank_after": found_blank, "watermark": "keep", "theme": "light"}}
