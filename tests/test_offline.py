"""Offline tests (no API key needed): python -m pytest -q tests"""
import os

from pptx import Presentation
from pptx.util import Emu

from qdeck.assemble import Fragment, assemble
from qdeck.config import Settings
from qdeck.deck import Q1, SLIDE_W, SLIDE_H

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def F(page, order, role, num=None, lang="en", cont_prev=False, cont_next=False, text="", conf=0.9):
    return Fragment(page_seq=page, page_ref=f"p{page + 1}", order=order, role=role,
                    label=None if num is None else f"{num}.", number=num, lang=lang, box=(0, 0, 1, 1),
                    crop_path="", text=text or f"question {num} {lang}", options=[], confidence=conf,
                    continues_previous=cont_prev, continues_on_next=cont_next)


def test_column_split_is_rejoined():
    frs = [F(0, 0, "question", 7), F(0, 1, "question", 8, cont_next=True),
           F(0, 2, "continuation", cont_prev=True, text="(a) 1 (b) 2 (c) 3 (d) 4"), F(0, 3, "question", 9)]
    qs, dropped, integ = assemble(frs)
    assert [q.number for q in qs] == [7, 8, 9]
    assert len(qs[1].parts) == 1 and "split_stitched" in qs[1].flags
    assert integ["ok"] and integ["expected"] == 3


def test_page_split_goes_to_same_language_question():
    # Hindi page then English page on one spread; the Hindi question continues on the next spread
    frs = [F(0, 0, "question", 51, "hi", cont_next=True), F(0, 1, "question", 51, "en"),
           F(1, 0, "continuation", lang="hi", cont_prev=True), F(1, 1, "question", 52, "hi"),
           F(1, 2, "question", 52, "en")]
    qs, dropped, integ = assemble(frs, lang_pref="en")
    assert [(q.number, q.lang) for q in qs] == [(51, "en"), (52, "en")]
    assert dropped["twins"] == 2 and dropped["orphans"] == 0


def test_solutions_clutter_and_duplicates_dropped():
    frs = [F(0, 0, "clutter"), F(0, 1, "question", 1), F(0, 2, "solution"), F(0, 3, "question", 2),
           F(1, 0, "question", 1), F(1, 1, "question", 2), F(1, 2, "question", 3)]
    qs, dropped, integ = assemble(frs)
    assert [q.number for q in qs] == [1, 2, 3]
    assert dropped["duplicates"] == 2 and dropped["solutions"] == 1


def test_numbering_restart_makes_new_section_and_gaps_are_reported():
    frs = [F(0, i, "question", n, text=f"section {i // 3} physics question number {n}" if i < 3 else f"chemistry item {n} about acids")
           for i, n in enumerate([1, 2, 4, 1, 2])]
    qs, _, integ = assemble(frs)
    assert len(qs) == 5 and integ["gaps"] == ["S1:3"] and len(integ["sections"]) == 2
    assert not integ["ok"]


def test_quadrant_geometry():
    left, top, w, h = Q1
    assert left >= SLIDE_W / 2 and left + w <= SLIDE_W and top >= 0 and top + h <= SLIDE_H / 2


def test_electrician_without_llm(tmp_path):
    pdf = os.path.join(ROOT, "samples", "electrician.pdf")
    if not os.path.exists(pdf):
        import pytest
        pytest.skip("sample not present")
    from qdeck.pipeline import run_job
    s = Settings()
    s.use_llm = False
    s.cache_dir = str(tmp_path / "cache")
    res = run_job([pdf], str(tmp_path / "out"), s)
    assert res["slides"] == 75 and res["integrity"]["ok"]
    prs = Presentation(res["pptx"])
    assert len(prs.slides) == 150                        # question + blank working slide
    assert abs(prs.slide_width - Emu(int(13.333 * 914400))) < 2000
    emu = 914400
    for i, slide in enumerate(prs.slides):
        if i % 2:
            assert len(slide.shapes) == 0                # blank slide for the teacher
            continue
        assert len(slide.shapes) == 1
        sh = slide.shapes[0]                             # picture inside the slide margins
        assert sh.left >= int(0.29 * emu) and sh.top >= int(0.29 * emu)
        assert sh.left + sh.width <= int((SLIDE_W - 0.29) * emu) and sh.top + sh.height <= int((SLIDE_H - 0.29) * emu)
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    # no cross-hair axes / quadrant lines anywhere: the master holds only its placeholders
    assert all(sh.shape_type == MSO_SHAPE_TYPE.PLACEHOLDER for sh in prs.slide_master.shapes)
    assert os.path.getsize(res["pptx"]) < 15 * 1024 * 1024
