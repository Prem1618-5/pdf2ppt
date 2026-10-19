"""Offline test for qdeck.pptx_import: python -m pytest -q tests/test_pptx_import.py"""
import io
import os

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from qdeck.pptx_import import import_pptx


def _png(w, h, color):
    buf = io.BytesIO()
    Image.new("RGB", (w, h), color).save(buf, "PNG")
    buf.seek(0)
    return buf


def _build(path):
    prs = Presentation()  # default template is 4:3 (10in x 7.5in)
    blank, title_content = prs.slide_layouts[6], prs.slide_layouts[1]

    s1 = prs.slides.add_slide(blank)
    s1.shapes.add_picture(_png(400, 200, "red"), Inches(1), Inches(1), Inches(4), Inches(2))
    s1.notes_slide.notes_text_frame.text = "Q 7. | source: x"

    prs.slides.add_slide(title_content)  # only empty placeholders -> blank working slide

    s3 = prs.slides.add_slide(blank)
    pic = s3.shapes.add_picture(_png(300, 300, "blue"), Inches(1), Inches(1), Inches(3), Inches(3))
    pic.crop_left = 0.1

    s4 = prs.slides.add_slide(blank)
    tb = s4.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(4))
    tb.text_frame.text = "Q.43 Who was...\n(a) A\n(b) B\n(c) C\n(d) D"

    s5 = prs.slides.add_slide(blank)  # pictures inside a group; the larger one wins
    grp = s5.shapes.add_group_shape()
    grp.shapes.add_picture(_png(50, 50, "green"), Inches(0), Inches(0), Inches(1), Inches(1))
    grp.shapes.add_picture(_png(640, 320, "yellow"), Inches(2), Inches(2), Inches(6), Inches(3))
    prs.save(path)


def test_import_pptx(tmp_path):
    src = str(tmp_path / "deck.pptx")
    out = str(tmp_path / "out")
    _build(src)
    proj = import_pptx(src, out)
    items, pages = proj["items"], proj["pages"]

    assert [it["label"] for it in items] == ["7.", "2", "Q.43", "4"]
    assert [it["number"] for it in items] == [7, 2, 43, 4]
    assert len({it["id"] for it in items}) == 4 and items[0]["id"] == "q0007"
    # 1 of 4 slides is followed by a blank: deck-wide setting off, that slide overrides it
    assert [it["blank_after"] for it in items] == [True, None, None, None]
    assert proj["settings"] == {"layout": "fill", "blank_after": False, "watermark": "keep", "theme": "light"}

    assert items[0]["render"] == "snapshot" and items[0]["parts"] == [{"page": 0, "box": [0.0, 0.0, 1.0, 1.0]}]
    x0, y0, x1, y1 = items[1]["parts"][0]["box"]
    assert abs(x0 - 0.1) < 1e-3 and (y0, x1, y1) == (0.0, 1.0, 1.0)

    q = items[2]
    assert q["render"] == "text" and q["parts"] == []
    assert q["stem"] == "Who was..." and q["options"] == ["A", "B", "C", "D"]

    assert items[3]["parts"][0]["page"] == 2
    assert [(p["seq"], p["ref"], p["w"], p["h"]) for p in pages] == [
        (0, "deck.pptx#s1", 400, 200), (1, "deck.pptx#s3", 300, 300), (2, "deck.pptx#s5", 640, 320)]
    for p in pages:
        assert os.path.isabs(p["path"]) and os.path.isfile(p["path"]) and p["path"].endswith(".png")
        assert p["kind"] == "image"
    assert os.path.basename(pages[0]["path"]) == "page_000.png"
