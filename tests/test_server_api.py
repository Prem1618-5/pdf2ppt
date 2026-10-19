"""Editor API end-to-end on the born-digital sample (no network, no keys)."""
import os
import time

import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PDF = os.path.join(ROOT, "samples", "electrician.pdf")


@pytest.fixture(scope="module")
def client(tmp_path_factory):
    if not os.path.exists(PDF):
        pytest.skip("sample not present")
    os.environ["JOBS_DIR"] = str(tmp_path_factory.mktemp("jobs"))
    os.environ["GEMINI_API_KEYS"] = ""
    import importlib
    import qdeck.server as server
    importlib.reload(server)
    from fastapi.testclient import TestClient
    with TestClient(server.app) as c:
        yield c


def _wait(c, jid):
    for _ in range(300):
        s = c.get(f"/api/jobs/{jid}").json()
        if s["state"] in ("done", "error"):
            return s
        time.sleep(0.5)
    raise AssertionError("timeout")


def test_full_edit_flow(client):
    c = client
    with open(PDF, "rb") as f:
        r = c.post("/api/jobs", files=[("files", ("electrician.pdf", f, "application/pdf"))],
                   data={"keys": "", "lang": "en", "render": "auto"})
    assert r.status_code == 200, r.text
    jid = r.json()["id"]
    assert _wait(c, jid)["state"] == "done"
    p = c.get(f"/api/jobs/{jid}/project").json()
    assert len(p["items"]) == 75 and p["settings"]["blank_after"] is True and p["settings"]["watermark"] == "remove"
    first = p["items"][0]
    assert c.get(f"/api/jobs/{jid}/pages/0.jpg").status_code == 200
    assert c.get(f"/api/jobs/{jid}/items/{first['id']}.png").headers["content-type"] == "image/png"
    # edit label + crop
    box = first["parts"][0]["box"]
    r = c.patch(f"/api/jobs/{jid}/items/{first['id']}",
                json={"label": "1.", "parts": [{"page": first["parts"][0]["page"], "box": [box[0], box[1], box[2], box[1] + 0.1]}]})
    assert r.status_code == 200 and r.json()["item"]["rev"] == first["rev"] + 1
    # delete, duplicate, merge, reorder, settings, undo
    ids = [x["id"] for x in p["items"]]
    after_delete = c.delete(f"/api/jobs/{jid}/items/{ids[5]}").json()["project"]
    assert len(after_delete["items"]) == 74
    assert after_delete["integrity"]["ok"] is False and after_delete["integrity"]["gaps"]   # live numbering check
    assert len(c.post(f"/api/jobs/{jid}/items/{ids[6]}/duplicate").json()["project"]["items"]) == 75
    proj = c.post(f"/api/jobs/{jid}/items/{ids[1]}/merge", json={"with": ids[2]}).json()["project"]
    assert len(proj["items"]) == 74
    new = [x["id"] for x in proj["items"]][::-1]
    assert [x["id"] for x in c.post(f"/api/jobs/{jid}/order", json={"ids": new}).json()["project"]["items"]] == new
    assert c.patch(f"/api/jobs/{jid}/settings", json={"blank_after": False}).json()["project"]["settings"]["blank_after"] is False
    undone = c.post(f"/api/jobs/{jid}/undo").json()["project"]
    assert undone["settings"]["blank_after"] is True
    # new question from a drawn box
    r = c.post(f"/api/jobs/{jid}/items", json={"after": None, "label": "76.", "parts": [{"page": 3, "box": [0.1, 0.1, 0.9, 0.3]}]})
    assert r.status_code == 200
    # export
    ex = c.post(f"/api/jobs/{jid}/export").json()
    from pptx import Presentation
    import io
    deck = Presentation(io.BytesIO(c.get(ex["url"]).content))
    n_items = len(c.get(f"/api/jobs/{jid}/project").json()["items"])
    assert len(deck.slides) == ex["slides"] == 2 * n_items     # blank slide after each question
    for i, slide in enumerate(deck.slides):
        assert (len(slide.shapes) == 0) == (i % 2 == 1)
    assert c.post("/api/keys/check", json={"keys": []}).json() == []
    assert c.delete(f"/api/jobs/{jid}").json()["deleted"]


def test_image_request_does_not_undo_concurrent_edit(client):
    """Regression: GET item png used to save a stale project copy over a fresh edit."""
    import threading
    c = client
    with open(PDF, "rb") as f:
        jid = c.post("/api/jobs", files=[("files", ("electrician.pdf", f, "application/pdf"))]).json()["id"]
    assert _wait(c, jid)["state"] == "done"
    items = c.get(f"/api/jobs/{jid}/project").json()["items"]
    a, b = items[10], items[20]
    box = a["parts"][0]["box"]
    t = threading.Thread(target=lambda: c.patch(f"/api/jobs/{jid}/items/{a['id']}", json={
        "label": "RACE", "parts": [{"page": a["parts"][0]["page"], "box": [box[0], box[1], box[2], box[1] + 0.05]}]}))
    t.start()
    for _ in range(5):
        c.get(f"/api/jobs/{jid}/items/{b['id']}.png")
    t.join()
    after = {x["id"]: x for x in c.get(f"/api/jobs/{jid}/project").json()["items"]}
    assert after[a["id"]]["label"] == "RACE" and after[a["id"]]["rev"] == a["rev"] + 1
